from __future__ import annotations

from attack_runner import AttackResult


def export_results_to_excel(results: list[AttackResult], filepath: str) -> None:
    """Write results to an Excel workbook with multiple sheets."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    green = PatternFill("solid", fgColor="D4EDDA")
    red = PatternFill("solid", fgColor="F8D7DA")
    orange = PatternFill("solid", fgColor="FFF3CD")
    purple = PatternFill("solid", fgColor="E2D9F3")
    header = PatternFill("solid", fgColor="343A40")

    bold_white = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    def style_header_row(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, ncols: int) -> None:
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = header
            cell.font = bold_white
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def fill_for_verdict(verdict: str) -> PatternFill | None:
        return {
            "SAFE": green,
            "VULNERABLE": red,
            "ERROR": orange,
            "DEMO": purple,
        }.get(verdict)

    def auto_col_width(ws: openpyxl.worksheet.worksheet.Worksheet, max_w: int = 60) -> None:
        for col_cells in ws.columns:
            length = max((len(str(c.value or "")) for c in col_cells), default=10)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 4, max_w)

    ws_all = wb.active
    ws_all.title = "All Results"

    headers = ["LLM", "Category", "Test Name", "Verdict", "Prompt", "Response", "Details", "Timestamp"]
    ws_all.append(headers)
    style_header_row(ws_all, 1, len(headers))

    for result in results:
        ws_all.append(
            [
                result.llm_name,
                result.attack_category,
                result.test_name,
                result.verdict,
                result.prompt,
                result.response,
                result.details,
                result.timestamp,
            ]
        )
        fill = fill_for_verdict(result.verdict)
        row_idx = ws_all.max_row
        for col in range(1, len(headers) + 1):
            cell = ws_all.cell(row=row_idx, column=col)
            cell.alignment = wrap
            cell.border = thin_border
            if fill:
                cell.fill = fill

    ws_all.row_dimensions[1].height = 20
    ws_all.freeze_panes = "A2"
    auto_col_width(ws_all)

    ws_cmp = wb.create_sheet("LLM Comparison")
    llm_names = list(dict.fromkeys(r.llm_name for r in results if r.llm_name != "N/A"))
    test_keys = list(dict.fromkeys(f"{r.attack_category} | {r.test_name}" for r in results if r.llm_name != "N/A"))

    if llm_names and test_keys:
        lookup: dict[tuple[str, str], str] = {}
        for result in results:
            if result.llm_name == "N/A":
                continue
            key = f"{result.attack_category} | {result.test_name}"
            lookup[(result.llm_name, key)] = result.verdict

        ws_cmp.append(["Test"] + llm_names)
        style_header_row(ws_cmp, 1, len(llm_names) + 1)

        for test_key in test_keys:
            ws_cmp.append([test_key] + [lookup.get((llm, test_key), "-") for llm in llm_names])
            row_idx = ws_cmp.max_row
            ws_cmp.cell(row=row_idx, column=1).font = bold
            ws_cmp.cell(row=row_idx, column=1).border = thin_border
            ws_cmp.cell(row=row_idx, column=1).alignment = wrap

            for col_idx, llm in enumerate(llm_names, start=2):
                verdict = lookup.get((llm, test_key), "-")
                cell = ws_cmp.cell(row=row_idx, column=col_idx)
                cell.value = verdict
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                fill = fill_for_verdict(verdict)
                if fill:
                    cell.fill = fill

        ws_cmp.freeze_panes = "B2"
        auto_col_width(ws_cmp, max_w=40)

    categories = list(dict.fromkeys(r.attack_category for r in results))
    for category in categories:
        category_results = [r for r in results if r.attack_category == category]
        sheet_name = category[:31].translate(str.maketrans(r"/\?*:[]-", "________"))
        ws_cat = wb.create_sheet(sheet_name)

        cat_headers = ["LLM", "Test Name", "Verdict", "Prompt (preview)", "Response (preview)", "Details"]
        ws_cat.append(cat_headers)
        style_header_row(ws_cat, 1, len(cat_headers))

        for result in category_results:
            ws_cat.append(
                [
                    result.llm_name,
                    result.test_name,
                    result.verdict,
                    result.prompt[:200],
                    result.response[:300],
                    result.details,
                ]
            )
            fill = fill_for_verdict(result.verdict)
            row_idx = ws_cat.max_row
            for col in range(1, len(cat_headers) + 1):
                cell = ws_cat.cell(row=row_idx, column=col)
                cell.alignment = wrap
                cell.border = thin_border
                if fill:
                    cell.fill = fill

        ws_cat.freeze_panes = "A2"
        auto_col_width(ws_cat)

    wb.save(filepath)

